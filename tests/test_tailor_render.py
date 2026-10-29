import pdfplumber
from docx import Document

from jobbot import ats
from jobbot.render import cv_text, render_all
from jobbot.tailor import _Bullet, _Role, _SkillGroup, _TailoredOutput, keyword_tailor, llm_tailor, validate_tailored


def _output(**kw):
    base = dict(
        headline="Senior Data Analyst | SQL & Power BI",
        summary="Data analyst with 7 years of experience in retail and fintech analytics.",
        skill_groups=[_SkillGroup(category="Analytics", skills=["SQL", "Power BI", "Python"])],
        roles=[
            _Role(experience_id="e0", bullets=[
                _Bullet(source_ids=["e0b0"], text="Built Power BI dashboards used by 120 store and category managers"),
                _Bullet(source_ids=["e0b2"], text="Built a Python forecasting model that reduced stock-outs by 18%"),
            ]),
            _Role(experience_id="e1", bullets=[
                _Bullet(source_ids=["e1b0"], text="Built SQL pipelines and Tableau dashboards for fraud KPIs"),
            ]),
        ],
        keywords_targeted=["SQL", "Power BI"], missing_keywords=["Looker"],
        cover_letter="Dear Hiring Team,\n\nI have 7 years of experience.\n\nKind regards,\nAlex Example",
    )
    base.update(kw)
    return _TailoredOutput(**base)


def test_valid_output_is_kept(profile):
    cv = validate_tailored(profile, _output())
    assert cv.warnings == []
    assert cv.experience[0].bullets[0].startswith("Built Power BI dashboards")
    assert cv.experience[0].company == "Example Retail Group" and cv.experience[1].start == "Jan 2018"
    assert list(cv.skills)[0] == "Analytics"
    # skills the model didn't place are kept, not lost
    assert "Tableau" in cv.skills["Additional Skills"]


def test_invented_numbers_skills_and_bullets_are_rejected(profile):
    out = _output(
        summary="Data analyst with 12 years of experience.",
        skill_groups=[_SkillGroup(category="Analytics", skills=["SQL", "Kubernetes"])],
        roles=[_Role(experience_id="e0", bullets=[
            _Bullet(source_ids=["e0b0"], text="Built dashboards used by 500 managers"),          # invented number
            _Bullet(source_ids=["e9b9"], text="Led a team of engineers"),                         # no such source
            _Bullet(source_ids=["e1b0"], text="Wrote SQL pipelines"),                             # source from another role
        ])],
        cover_letter="I increased revenue by 300%.",
    )
    cv = validate_tailored(profile, out)
    text = cv_text(cv, profile)
    assert "500" not in text and "Led a team" not in text and "Kubernetes" not in text
    assert "120 store" in cv.experience[0].bullets[0]  # reverted to original wording
    assert cv.summary == profile.summary
    assert cv.cover_letter == ""
    # a role the model dropped falls back to the original bullets
    assert cv.experience[1].bullets == profile.experience[1].bullets[:4]
    assert len(cv.warnings) >= 5


def test_llm_tailor_uses_parsed_output(profile, job):
    class FakeLLM:
        def parse(self, **kwargs):
            assert "Senior Data Analyst" in kwargs["prompt"] and "<master_cv>" in kwargs["cached_context"]
            return _output()

    cv = llm_tailor(FakeLLM(), profile, job, ["SQL", "Power BI"])
    assert cv.method == "llm" and cv.cover_letter.startswith("Dear Hiring Team")


def test_keyword_tailor_puts_relevant_first(profile, job):
    kws = ats.extract_keywords(job.description, profile.all_skills(), job.title)
    cv = keyword_tailor(profile, job, kws)
    assert cv.skills["Analytics"][0] in ("SQL", "Python", "Power BI", "Tableau")
    assert job.company in cv.cover_letter and profile.name in cv.cover_letter
    assert "[" not in cv.cover_letter


def test_render_pdf_and_docx_are_ats_readable(profile, job, tmp_path):
    cv = keyword_tailor(profile, job, ["SQL", "Python"])
    files = render_all(cv, profile, job, tmp_path, ["pdf", "docx"])
    assert {"cv_pdf", "cv_docx", "cover_letter_pdf", "cover_letter_txt"} <= set(files)
    with pdfplumber.open(files["cv_pdf"]) as pdf:
        pdf_text = "\n".join(p.extract_text() for p in pdf.pages)
    docx_text = "\n".join(p.text for p in Document(str(files["cv_docx"])).paragraphs)
    for text in (pdf_text, docx_text):
        assert profile.personal["email"] in text
        assert "PROFESSIONAL EXPERIENCE" in text and "Example Retail Group" in text
        assert "Visa: UAE residence visa" in text
    assert Document(str(files["cv_docx"])).tables == []  # no tables: ATS-safe layout
    assert ats.lint_cv_text(cv_text(cv, profile), profile.personal["email"], profile.personal["phone"]) == []
    assert files["cv_pdf"].name == "Alex_Example_Acme_Payments_Senior_Data_Analyst_CV.pdf"
