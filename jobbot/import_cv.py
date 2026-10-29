"""Turn an existing CV (PDF/DOCX/TXT) into a profile.yaml draft using Claude."""

from __future__ import annotations

from pathlib import Path

import yaml
from pydantic import BaseModel, Field


class _Personal(BaseModel):
    name: str
    email: str
    phone: str
    location: str = Field(description="City, Country - e.g. 'Dubai, UAE'")
    linkedin: str
    website: str
    nationality: str
    visa_status: str
    notice_period: str
    driving_license: str


class _SkillGroup(BaseModel):
    category: str
    skills: list[str]


class _Experience(BaseModel):
    company: str
    title: str
    location: str
    start: str = Field(description="'Mon YYYY', e.g. 'Mar 2021'")
    end: str = Field(description="'Mon YYYY' or 'Present'")
    bullets: list[str]


class _Education(BaseModel):
    institution: str
    degree: str
    field: str
    location: str
    year: str


class _Cert(BaseModel):
    name: str
    issuer: str
    year: str


class _ProfileModel(BaseModel):
    personal: _Personal
    headline: str
    summary: str
    skills: list[_SkillGroup]
    languages: list[str]
    experience: list[_Experience]
    education: list[_Education]
    certifications: list[_Cert]


SYSTEM = """You convert a CV into structured data. Copy facts exactly as written - do not improve, embellish,
infer or add anything. Use an empty string for any field the CV does not state. Keep every bullet point
(lightly cleaned of stray symbols), in the original order. Group skills into sensible categories using the
skill names as written in the CV."""


def read_cv_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            import pdfplumber
        except ImportError as exc:
            raise RuntimeError("reading PDFs needs pdfplumber: pip install pdfplumber") from exc
        with pdfplumber.open(str(path)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


def import_cv(llm, cv_path: Path, out_path: Path) -> Path:
    text = read_cv_text(cv_path)
    if len(text.strip()) < 200:
        raise RuntimeError("could not read much text from that CV - is it a scanned image? Try a DOCX version.")
    model: _ProfileModel = llm.parse(system=SYSTEM, prompt=f"<cv>\n{text}\n</cv>", schema=_ProfileModel)
    data = model.model_dump()
    data["personal"] = {k: v for k, v in data["personal"].items() if v}
    data["skills"] = {g["category"]: g["skills"] for g in data.pop("skills")}
    for group in ("education", "certifications"):
        data[group] = [{k: v for k, v in item.items() if v} for item in data[group]]
    data["application_answers"] = DEFAULT_ANSWERS
    header = ("# Generated from your CV by `jobbot import-cv`. Review every line: this file is the only\n"
              "# source of facts for your tailored CVs. Fill in application_answers at the bottom.\n")
    out_path.write_text(header + yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=110),
                        encoding="utf-8")
    return out_path


DEFAULT_ANSWERS = [
    {"match": "notice period", "answer": "1 month"},
    {"match": "salary|compensation|expected (pay|package)", "answer": "Negotiable"},
    {"match": "sponsor", "answer": "No"},
    {"match": "authori[sz]ed to work|right to work|eligib\\w* to work|work permit", "answer": "Yes"},
    {"match": "visa", "answer": "UAE residence visa"},
    {"match": "relocat", "answer": "Yes"},
    {"match": "how did you hear|where did you (hear|find)", "answer": "Company careers page"},
]
