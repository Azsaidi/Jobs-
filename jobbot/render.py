"""Render a tailored CV and cover letter as ATS-friendly DOCX and PDF.

ATS-friendly means: one column, no tables, text boxes, images or icons, contact details in
the body (not the page header), standard section headings, real text (not images), and a
common font.
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from .models import Job, TailoredCV
from .profile import Profile
from .util import slugify

FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]


def contact_line(profile: Profile) -> str:
    p = profile.personal
    items = [p.get("location"), p.get("phone"), p.get("email"), p.get("linkedin"), p.get("website")]
    return " | ".join(str(x) for x in items if x)


def personal_details_line(profile: Profile) -> str:
    """Nationality / visa / notice period - customary on UAE CVs and often an ATS knock-out filter."""
    p = profile.personal
    items = []
    if p.get("nationality"):
        items.append(f"Nationality: {p['nationality']}")
    if p.get("visa_status"):
        items.append(f"Visa: {p['visa_status']}")
    if p.get("notice_period"):
        items.append(f"Notice period: {p['notice_period']}")
    if p.get("driving_license"):
        items.append(f"Driving licence: {p['driving_license']}")
    return " | ".join(items)


def _edu_line(e: dict) -> tuple[str, str]:
    degree = ", ".join(str(x) for x in (e.get("degree"), e.get("field")) if x)
    where = " - ".join(str(x) for x in (e.get("institution"), e.get("location")) if x)
    when = e.get("year") or " - ".join(str(x) for x in (e.get("start"), e.get("end")) if x)
    head = " | ".join(x for x in (degree, where) if x)
    return head, str(when or "")


def _cert_line(c: dict | str) -> str:
    if isinstance(c, str):
        return c
    return " - ".join(str(x) for x in (c.get("name"), c.get("issuer"), c.get("year")) if x)


def sections(cv: TailoredCV, profile: Profile) -> list[tuple[str, list]]:
    """Format-independent CV structure: [(heading, [items])]."""
    out: list[tuple[str, list]] = [("Professional Summary", [("para", cv.summary)])]
    out.append(("Core Skills", [("para", f"{cat}: {', '.join(items)}") for cat, items in cv.skills.items() if items]))
    exp_items = []
    for e in cv.experience:
        exp_items.append(("role", f"{e.title} | {e.company}",
                          " | ".join(x for x in (e.location, f"{e.start} - {e.end}") if x), e.bullets))
    out.append(("Professional Experience", exp_items))
    if profile.education:
        out.append(("Education", [("role", *_edu_line(e), [str(d) for d in (e.get("details") or [])])
                                  for e in profile.education]))
    if profile.certifications:
        out.append(("Certifications", [("bullet", _cert_line(c)) for c in profile.certifications]))
    if profile.projects:
        out.append(("Projects", [("role", p.get("name", ""), p.get("description", ""), list(p.get("bullets") or []))
                                 for p in profile.projects]))
    if profile.languages:
        out.append(("Languages", [("para", ", ".join(profile.languages))]))
    return out


def cv_text(cv: TailoredCV, profile: Profile, personal_details: bool = True) -> str:
    """Plain-text version (what an ATS parser 'sees'); used for keyword scoring."""
    lines = [profile.name, cv.headline, contact_line(profile)]
    if personal_details and personal_details_line(profile):
        lines.append(personal_details_line(profile))
    for heading, items in sections(cv, profile):
        lines += ["", heading.upper()]
        for item in items:
            if item[0] == "role":
                lines += [item[1], item[2], *[f"- {b}" for b in item[3]]]
            else:
                lines.append(("- " if item[0] == "bullet" else "") + item[1])
    return "\n".join(x for x in lines if x is not None)


def output_basename(profile: Profile, job: Job) -> str:
    return f"{slugify(profile.name, 30)}_{slugify(job.company, 25)}_{slugify(job.title, 40)}"


# ---- DOCX -------------------------------------------------------------------

def render_docx(cv: TailoredCV, profile: Profile, path: Path, personal_details: bool = True) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    def para(text="", bold=False, size=None, align=None, space_after=None, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=11.5, space_after=3, color=(0x1F, 0x3A, 0x5F))
        p.paragraph_format.space_before = Pt(9)
        # bottom rule as a paragraph border (plain text for ATS, unlike a drawn line/shape)
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        p._p.get_or_add_pPr().append(pbdr)

    para(profile.name, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(cv.headline, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    para(contact_line(profile), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    if personal_details and personal_details_line(profile):
        para(personal_details_line(profile), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    for title, items in sections(cv, profile):
        heading(title)
        for item in items:
            if item[0] == "para":
                para(item[1])
            elif item[0] == "bullet":
                doc.add_paragraph(item[1], style="List Bullet")
            else:
                _, head, sub, bullets = item
                p = para(head, bold=True, space_after=0)
                p.paragraph_format.space_before = Pt(5)
                if sub:
                    para(sub, size=9.5, space_after=1, color=(0x55, 0x55, 0x55))
                for b in bullets:
                    bp = doc.add_paragraph(b, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(1)

    core = doc.core_properties
    core.author = profile.name
    core.title = f"{profile.name} - CV"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# ---- PDF --------------------------------------------------------------------

_fonts: tuple[str, str] | None = None


def _pdf_fonts() -> tuple[str, str]:
    """Embed a TrueType font when available (full Unicode, better text extraction), else Helvetica."""
    global _fonts
    if _fonts:
        return _fonts
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    _fonts = ("Helvetica", "Helvetica-Bold")
    for regular, bold in FONT_CANDIDATES:
        if Path(regular).is_file() and Path(bold).is_file():
            try:
                pdfmetrics.registerFont(TTFont("CVSans", regular))
                pdfmetrics.registerFont(TTFont("CVSans-Bold", bold))
                pdfmetrics.registerFontFamily("CVSans", normal="CVSans", bold="CVSans-Bold")
                _fonts = ("CVSans", "CVSans-Bold")
                break
            except Exception:  # noqa: BLE001 - unusable font file, try the next one
                continue
    return _fonts


def _styles():
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle

    regular, bold = _pdf_fonts()
    base = ParagraphStyle("base", fontName=regular, fontSize=10, leading=13)
    return {
        "name": ParagraphStyle("name", base, fontName=bold, fontSize=18, leading=22, alignment=TA_CENTER),
        "headline": ParagraphStyle("headline", base, fontSize=11.5, leading=15, alignment=TA_CENTER),
        "contact": ParagraphStyle("contact", base, fontSize=9, leading=12, alignment=TA_CENTER),
        "heading": ParagraphStyle("heading", base, fontName=bold, fontSize=11, leading=14, spaceBefore=9,
                                  spaceAfter=2, textColor=HexColor("#1F3A5F")),
        "body": base,
        "role": ParagraphStyle("role", base, fontName=bold, spaceBefore=5),
        "sub": ParagraphStyle("sub", base, fontSize=9, leading=12, textColor=HexColor("#555555")),
        "bullet": ParagraphStyle("bullet", base, leftIndent=12, bulletIndent=2, spaceAfter=1),
    }


def render_pdf(cv: TailoredCV, profile: Profile, path: Path, personal_details: bool = True) -> Path:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

    st = _styles()
    story = [Paragraph(escape(profile.name), st["name"]), Paragraph(escape(cv.headline), st["headline"]),
             Paragraph(escape(contact_line(profile)), st["contact"])]
    if personal_details and personal_details_line(profile):
        story.append(Paragraph(escape(personal_details_line(profile)), st["contact"]))
    for title, items in sections(cv, profile):
        story.append(Paragraph(escape(title.upper()), st["heading"]))
        story.append(HRFlowable(width="100%", thickness=0.6, color=HexColor("#1F3A5F"), spaceAfter=3))
        for item in items:
            if item[0] == "para":
                story.append(Paragraph(escape(item[1]), st["body"]))
            elif item[0] == "bullet":
                story.append(Paragraph(escape(item[1]), st["bullet"], bulletText="•"))
            else:
                _, head, sub, bullets = item
                story.append(Paragraph(escape(head), st["role"]))
                if sub:
                    story.append(Paragraph(escape(sub), st["sub"]))
                story += [Paragraph(escape(b), st["bullet"], bulletText="•") for b in bullets]
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=f"{profile.name} - CV",
                            author=profile.name)
    doc.build(story)
    return path


def render_cover_letter_pdf(text: str, profile: Profile, job: Job, path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    st = _styles()
    story = [Paragraph(escape(profile.name), st["name"]), Paragraph(escape(contact_line(profile)), st["contact"]),
             Spacer(1, 10 * mm), Paragraph(escape(f"Re: {job.title} - {job.company}"), st["role"]), Spacer(1, 4 * mm)]
    for block in text.split("\n\n"):
        story.append(Paragraph(escape(block).replace("\n", "<br/>"), st["body"]))
        story.append(Spacer(1, 3 * mm))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title=f"{profile.name} - Cover letter", author=profile.name).build(story)
    return path


def render_all(cv: TailoredCV, profile: Profile, job: Job, out_dir: Path, formats: list[str],
               personal_details: bool = True) -> dict[str, Path]:
    base = output_basename(profile, job)
    out_dir.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}
    if "pdf" in formats:
        files["cv_pdf"] = render_pdf(cv, profile, out_dir / f"{base}_CV.pdf", personal_details)
    if "docx" in formats:
        files["cv_docx"] = render_docx(cv, profile, out_dir / f"{base}_CV.docx", personal_details)
    if cv.cover_letter:
        txt = out_dir / f"{base}_Cover_Letter.txt"
        txt.write_text(cv.cover_letter, encoding="utf-8")
        files["cover_letter_txt"] = txt
        files["cover_letter_pdf"] = render_cover_letter_pdf(cv.cover_letter, profile, job,
                                                            out_dir / f"{base}_Cover_Letter.pdf")
    return files
